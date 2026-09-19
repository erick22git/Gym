"""
Servicio HTTP local que empaqueta el pipeline validado:
  Prueba A (reconstruccion posicional con PaddleOCR)
  + umbral de fusion extrema (>5 lineas crudas -> revision humana directa)
  + Prueba B (LLM qwen2.5:3b via Ollama, formato compacto, temperature=0)
  + verificacion de cobertura (tolerante a simbolos sueltos sin valor informativo)

Uso:
    venv-ocr\\Scripts\\python.exe scripts\\ocr_service.py

Expone:
    POST http://localhost:8420/procesar-imagen
        - multipart/form-data con campo "imagen" (archivo), o
        - JSON {"imagen_base64": "..."}
    POST http://localhost:8420/procesar-excel  - multipart, campo "archivo" (.xlsx)
    POST http://localhost:8420/procesar-word   - multipart, campo "archivo" (.docx)
    POST http://localhost:8420/procesar-pdf    - multipart, campo "archivo" (.pdf)
    POST http://localhost:8420/asistente-texto - JSON {"texto": "..."} -> clasifica la
        intencion (importar / pregunta_sistema / charla_general) y responde acorde,
        usando scripts/conocimiento_sistema.md como base para preguntas del sistema.

    GET  http://localhost:8420/health

Excel/Word(tablas)/PDF(tablas): texto ya estructurado, va directo a
"resueltos" sin pasar por el LLM ni el detector de ambiguedad (esos existen
para el desorden del OCR, que aca no aplica). Word sin tablas (texto libre)
y PDF con texto pero sin tabla SI pasan por el mismo Agente Estructurador
(LLM) que usan las imagenes. PDF sin texto extraible (escaneado) se
convierte a imagen pagina por pagina y reusa el pipeline de OCR completo.
"""
import base64
import io
import json
import re
import statistics
import time
import urllib.request
import uuid
from pathlib import Path

from flask import Flask, request, jsonify
from paddleocr import PaddleOCR

PORT = 8420
EXTREME_FUSION_THRESHOLD = 5
RATIO_THRESHOLD = 4.0
OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = "qwen2.5:3b"

app = Flask(__name__)


@app.after_request
def _add_cors_headers(resp):
    # Electron/Chromium bloquea la respuesta de fetch() a un puerto distinto
    # (CORS) si el servidor no manda estos headers, aunque la llamada en si
    # funcione bien server-side. Servicio 100% local (localhost:8420), sin
    # datos sensibles de terceros, asi que "*" es aceptable acá.
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    resp.headers["Access-Control-Allow-Headers"] = "Content-Type"
    return resp


@app.route("/procesar-imagen", methods=["OPTIONS"])
def _procesar_imagen_preflight():
    return ("", 204)

print("Cargando PaddleOCR (esto tarda una vez al iniciar el servicio)...")
_ocr = PaddleOCR(
    use_doc_orientation_classify=False,
    use_doc_unwarping=False,
    use_textline_orientation=True,
    lang="es",
)
print("PaddleOCR listo.")


# ---------------------------------------------------------------------------
# Prueba A: OCR + reconstruccion posicional
# ---------------------------------------------------------------------------

def cluster_1d(values, gap_factor=2.5, min_gap_px=15):
    if not values:
        return []
    pairs = sorted(range(len(values)), key=lambda i: values[i])
    sorted_vals = [values[i] for i in pairs]
    gaps = [sorted_vals[i + 1] - sorted_vals[i] for i in range(len(sorted_vals) - 1)]
    if gaps:
        median_gap = statistics.median(gaps) if statistics.median(gaps) > 0 else 1
        threshold = max(median_gap * gap_factor, min_gap_px)
    else:
        threshold = min_gap_px
    clusters, current = [], [pairs[0]]
    for k in range(1, len(sorted_vals)):
        if sorted_vals[k] - sorted_vals[k - 1] > threshold:
            clusters.append(current)
            current = []
        current.append(pairs[k])
    clusters.append(current)
    return clusters


# Conectores cortos conocidos que PaddleOCR a veces devuelve pegados al
# numero siguiente dentro de UN SOLO fragmento detectado (ej. "DE6" en vez
# de "DE 6") — separarlos con un espacio es seguro porque son palabras
# completas conocidas, no arriesga romper un codigo alfanumerico real de
# producto (ej. "SKU123" no empieza con ninguna de estas palabras).
CONECTORES_CORTOS_PEGABLES = {
    "DE", "DEL", "CON", "SIN", "LA", "EL", "LOS", "LAS", "UN", "UNA",
    "Y", "A", "EN", "PARA", "POR", "AL",
}


def limpiar_espaciado_ocr(texto):
    def _separar(m):
        palabra = m.group(1)
        if palabra.upper() in CONECTORES_CORTOS_PEGABLES:
            return f"{palabra} {m.group(2)}"
        return m.group(0)
    # letra(s)+numero(s) pegados, ej. "DE6" -> grupo1="DE", grupo2="6"
    return re.sub(r"\b([A-Za-zÁÉÍÓÚÜáéíóúüÑñ]+)(\d+)\b", _separar, texto)


def run_ocr_and_reconstruct(image_path):
    result = _ocr.predict(str(image_path))
    res = result[0]
    texts, scores, boxes = res["rec_texts"], res["rec_scores"], res["rec_boxes"]

    items = []
    for text, score, box in zip(texts, scores, boxes):
        x1, y1, x2, y2 = [float(v) for v in box]
        items.append({"text": limpiar_espaciado_ocr(text), "score": float(score), "x1": x1, "y1": y1,
                      "x2": x2, "y2": y2, "yc": (y1 + y2) / 2, "h": y2 - y1})

    if not items:
        return {"n_lines": 0, "n_rows": 0, "table": [], "rows_raw": []}

    heights = [it["h"] for it in items]
    median_h = statistics.median(heights) if heights else 20

    items_sorted_y = sorted(items, key=lambda it: it["yc"])
    row_threshold = median_h * 0.7
    rows, current_row = [], [items_sorted_y[0]]
    for it in items_sorted_y[1:]:
        if it["yc"] - current_row[-1]["yc"] > row_threshold:
            rows.append(current_row)
            current_row = []
        current_row.append(it)
    rows.append(current_row)

    all_x1 = [it["x1"] for it in items]
    x_clusters_idx = cluster_1d(all_x1, gap_factor=2.0, min_gap_px=25)
    col_ranges = sorted(
        ((min(all_x1[i] for i in c), max(all_x1[i] for i in c)) for c in x_clusters_idx),
        key=lambda r: r[0],
    )

    def col_index(x1):
        best_i, best_d = 0, float("inf")
        for i, (lo, hi) in enumerate(col_ranges):
            d = abs(x1 - (lo + hi) / 2)
            if d < best_d:
                best_d, best_i = d, i
        return best_i

    table = []
    for row in rows:
        row_sorted = sorted(row, key=lambda it: it["x1"])
        cells = {}
        for it in row_sorted:
            cells.setdefault(col_index(it["x1"]), []).append(it["text"])
        table.append([" ".join(cells.get(c, [])) for c in range(len(col_ranges))])

    rows_raw = [
        [{"text": it["text"], "score": it["score"], "x1": it["x1"], "y1": it["y1"],
          "x2": it["x2"], "y2": it["y2"]} for it in sorted(row, key=lambda it: it["x1"])]
        for row in rows
    ]

    return {"n_lines": len(items), "n_rows": len(rows), "n_cols": len(col_ranges),
            "table": table, "rows_raw": rows_raw}


def pick_name_and_qty_columns(table):
    """Heuristica: la columna 'nombre' es la que mas texto alfabetico acumula;
    la columna 'precio' es la que mas '$' tiene; 'cantidad' es la mejor
    columna numerica restante (sin '$')."""
    if not table or not table[0]:
        return 0, None
    n_cols = len(table[0])
    alpha_score = [0] * n_cols
    dollar_count = [0] * n_cols
    numeric_count = [0] * n_cols
    for row in table:
        for c, cell in enumerate(row):
            alpha_score[c] += len(re.findall(r"[A-Za-z]", cell))
            dollar_count[c] += cell.count("$")
            if re.fullmatch(r"\s*\d+(\.\d+)?\s*", cell):
                numeric_count[c] += 1

    name_col = max(range(n_cols), key=lambda c: alpha_score[c])
    price_col = max(range(n_cols), key=lambda c: dollar_count[c]) if any(dollar_count) else None

    candidates = [c for c in range(n_cols) if c != name_col and c != price_col]
    qty_col = max(candidates, key=lambda c: numeric_count[c]) if candidates else None
    return name_col, qty_col


# ---------------------------------------------------------------------------
# Deteccion de ambiguedad (reglas validadas en las pruebas)
# ---------------------------------------------------------------------------

def is_ambiguous(row, name_word_limit=7):
    joined = " ".join(row)
    n_dollars = joined.count("$")
    has_digit = bool(re.search(r"\d", joined))
    if n_dollars > 1:
        return True, "multiples $ en la fila"
    if not has_digit and joined.strip() and "DESCRIPCION" not in joined and "CANTIDA" not in joined:
        return True, "sin digitos, probable nota marginal"
    for cell in row:
        numbers = re.findall(r"(?<!\S)\d+(?:\.\d+)?(?!\S)", cell)
        if len(numbers) >= 2:
            return True, f"columna con {len(numbers)} numeros sueltos"
    for cell in row:
        if re.search(r"[A-Za-z]{3,}", cell) and len(cell.split()) > name_word_limit:
            return True, "nombre sospechosamente largo (posible fusion)"
    if re.search(r"\$[^$]*[A-Za-z]{3,}[^$]*\$", joined):
        return True, "patron $...texto...$ (precios con producto entre medio)"
    return False, None


def is_informative(text):
    return bool(re.search(r"[A-Za-z0-9]", text))


# ---------------------------------------------------------------------------
# Prueba B: LLM con formato compacto + verificacion de cobertura
# ---------------------------------------------------------------------------

def call_ollama(prompt, model=OLLAMA_MODEL):
    body = json.dumps({
        "model": model, "prompt": prompt, "stream": False,
        "options": {"temperature": 0},
    }).encode("utf-8")
    req = urllib.request.Request(OLLAMA_URL, data=body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=600) as resp:
        return json.loads(resp.read().decode("utf-8"))


def extract_json_array(text):
    m = re.search(r"\[.*\]", text, re.DOTALL)
    if not m:
        return []
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError:
        return []


def build_single_group_prompt(raw_lines):
    n_raw = len(raw_lines)
    lines_repr = "\n".join(
        f'  [{li}] "{it["text"]}" @x={int(it["x1"])},y={int(it["y1"])}'
        for li, it in enumerate(raw_lines)
    )
    return f"""Estos son fragmentos de texto (de OCR o de un documento) de una lista de
inventario. Podrian ser un solo producto o varios productos mezclados por error. Formato de
cada fragmento: "texto" @x=columna,y=fila (posicion aproximada; en texto plano "y" solo indica
el orden de las lineas, de arriba hacia abajo).

Fragmentos (indices de 0 a {n_raw-1}):
{lines_repr}

Identifica el/los producto(s) que representan estos fragmentos, con su cantidad si se puede
determinar. Para CADA producto que devuelvas, DEBES incluir "lineas_origen": la lista de
TODOS los indices de fragmentos que usaste para llegar a ese producto (nombre y cantidad),
incluyendo los que sean solo numeros o "$" si los usaste. Si un fragmento no lo pudiste usar
para nada, igual debes decidir a que producto pertenece o marcarlo aparte con "nombre":
"SIN_IDENTIFICAR". No inventes cantidades sin respaldo en un fragmento.

IMPORTANTE sobre el campo "nombre": debe ser la concatenacion LITERAL del texto de los
fragmentos de nombre citados en "lineas_origen", tal cual aparecen. NO resumas, NO acortes,
NO elimines palabras para que suene mas corto o mas prolijo (ej. si el fragmento dice "Coca
cola lightde 12 pzas 355 ml", el nombre debe incluir esa frase completa, no solo "Coca cola
light").

Devuelve SOLO un JSON: una lista de objetos con "nombre", "cantidad" (numero o null) y
"lineas_origen" (lista de numeros). No expliques nada, solo el JSON.
"""


def reconstruir_nombre_literal(lineas_origen, raw_lines):
    """Concatena literalmente el texto de las lineas citadas en
    lineas_origen que no sean puramente numericas/simbolicas (esas son la
    cantidad o ruido, no el nombre), en el orden en que aparecen. Sirve
    como respaldo cuando el LLM resume/acorta el nombre en vez de citarlo
    tal cual (ver resolver_grupo_llm)."""
    indices = sorted(i for i in (lineas_origen or []) if isinstance(i, int) and 0 <= i < len(raw_lines))
    partes = []
    for i in indices:
        texto = raw_lines[i]["text"].strip()
        if not texto:
            continue
        if re.fullmatch(r"\d+(\.\d+)?\s*\$?", texto) or re.fullmatch(r"[\W_]+", texto):
            continue  # linea puramente numerica o de simbolos -> no es el nombre
        partes.append(texto)
    return limpiar_espaciado_ocr(" ".join(partes).strip())


def limpiar_clausula_cantidad_del_nombre(nombre, cantidad):
    """Si el nombre trae pegada al final una clausula que solo repite la
    cantidad ya extraida por separado (ej. "Paracetamol 500mg - 120
    unidades", "Ibuprofeno 400mg, cantidad: 80", "Amoxicilina 500mg x60"),
    la recorta — es redundante, no es parte del nombre del producto.
    Corre SIEMPRE (venga el nombre del LLM o de reconstruir_nombre_literal)
    para no depender de que el LLM decida bien cuando resumir y cuando no."""
    if not nombre or cantidad is None:
        return nombre
    cantidad_str = str(int(cantidad)) if isinstance(cantidad, (int, float)) and float(cantidad).is_integer() else str(cantidad)
    patron = re.compile(
        rf"\s*[-,x×]\s*(cantidad\s*[:=]\s*)?{re.escape(cantidad_str)}\s*(unidades?|uds?\.?|piezas?|pzs?\.?)?\s*$",
        re.IGNORECASE,
    )
    limpio = patron.sub("", nombre).strip()
    return limpio if limpio else nombre


def resolver_grupo_llm(raw_lines, id_label):
    """Nucleo reutilizable de la Prueba B: una llamada al LLM para UN grupo de
    lineas (con o sin coordenadas reales), + verificacion de cobertura +
    Filtro 1 (nombre valido) + Filtro 2 (sin reuso sospechoso de linea) +
    cantidad real no nula. Usado tanto para grupos ambiguos de OCR (imagenes)
    como para texto libre de Word/PDF (ver structure_free_text_with_llm)."""
    n_raw = len(raw_lines)
    prompt = build_single_group_prompt(raw_lines)
    result = call_ollama(prompt)
    items = extract_json_array(result.get("response", ""))
    if isinstance(items, dict):
        items = [items]

    resueltos, requiere_revision = [], []

    covered = set()
    for it in items:
        for idx in it.get("lineas_origen", []) or []:
            covered.add(idx)
    expected = {idx for idx in range(n_raw) if is_informative(raw_lines[idx]["text"])}
    cobertura_ok = not (expected - covered) and bool(items)

    if not cobertura_ok:
        first = items[0] if items else {}
        requiere_revision.append({
            "id": id_label,
            "texto_ocr_crudo": [it["text"] for it in raw_lines],
            "nombre_sugerido": first.get("nombre"),
            "cantidad_sugerida": first.get("cantidad"),
            "motivo": "cobertura_incompleta",
        })
        return resueltos, requiere_revision

    for it in items:
        nombre = it.get("nombre")
        cantidad = it.get("cantidad")

        # Si el LLM resumio/acorto el nombre en vez de citarlo literal (ej.
        # "Coca cola light" cuando el fragmento citado decia "Coca cola
        # lightde 12 pzas 355 ml"), no confiamos en el resumen: reconstruimos
        # el nombre literal a partir de las mismas lineas_origen que el
        # modelo ya declaro haber usado, y si es notablemente mas largo (el
        # nombre del modelo es un prefijo/subconjunto del texto real), lo
        # reemplazamos. Independientemente de cual de los dos "gane", despues
        # se recorta cualquier clausula final que solo repita la cantidad
        # (ver limpiar_clausula_cantidad_del_nombre) — asi no dependemos de
        # que el LLM decida bien cuando una clausula de cantidad es
        # redundante y cuando el texto es informacion real del producto.
        nombre_reconstruido = reconstruir_nombre_literal(it.get("lineas_origen"), raw_lines)
        if nombre_reconstruido:
            nombre_norm = (nombre or "").strip().lower()
            reconstruido_norm = nombre_reconstruido.lower()
            if not nombre_norm or (reconstruido_norm.startswith(nombre_norm) and len(reconstruido_norm) > len(nombre_norm) * 1.2):
                nombre = nombre_reconstruido
        nombre = limpiar_clausula_cantidad_del_nombre(nombre, cantidad)

        ok1, _ = filtro1_nombre_valido(nombre)
        if not ok1:
            requiere_revision.append({
                "id": id_label,
                "texto_ocr_crudo": [it["text"] for it in raw_lines],
                "nombre_sugerido": nombre,
                "cantidad_sugerida": cantidad,
                "motivo": "nombre_invalido",
            })
            continue

        ok2, _ = filtro2_no_reuso_sospechoso(it, raw_lines)
        if not ok2:
            requiere_revision.append({
                "id": id_label,
                "texto_ocr_crudo": [it["text"] for it in raw_lines],
                "nombre_sugerido": nombre,
                "cantidad_sugerida": cantidad,
                "motivo": "origen_sospechoso",
            })
            continue

        if cantidad is None or not isinstance(cantidad, (int, float)):
            requiere_revision.append({
                "id": id_label,
                "texto_ocr_crudo": [it["text"] for it in raw_lines],
                "nombre_sugerido": nombre,
                "cantidad_sugerida": cantidad,
                "motivo": "sin_cantidad",
            })
            continue

        resueltos.append({"nombre": nombre, "cantidad": cantidad, "confianza": "alta"})

    return resueltos, requiere_revision


def filtro1_nombre_valido(nombre):
    """Rechaza: solo simbolos/puntuacion, numero puro, generico tipo
    SIN_IDENTIFICAR, o nombres de 1-2 caracteres (ej. "Q" — ningun producto
    real de inventario se llama con una sola letra, es casi siempre un
    fragmento suelto de OCR). Se aplica a TODAS las filas, no solo a las
    que pasan por el LLM — un nombre invalido es invalido venga de donde
    venga (ver CAMBIO 1: antes solo corria en resolver_grupo_llm, y por eso
    "Q" se colaba directo a resueltos desde la reconstruccion posicional)."""
    if not nombre:
        return False, "nombre vacio"
    n = str(nombre).strip()
    if len(n) <= 2:
        return False, "nombre_invalido"
    if re.fullmatch(r"[\W_]+", n):
        return False, "nombre_invalido"
    if re.fullmatch(r"\d+(\.\d+)?", n):
        return False, "nombre_invalido"
    if n.upper() in ("SIN_IDENTIFICAR", "N/A", "DESCONOCIDO", "SIN IDENTIFICAR"):
        return False, "nombre_invalido"
    return True, None


def filtro2_no_reuso_sospechoso(item, raw_lines):
    """Si cantidad no es None, debe estar respaldada por una linea PURAMENTE
    numerica citada en lineas_origen (no un numero embebido dentro del texto
    del nombre, ej. 'de 4 pzs' o 'DE 50 PZAS').

    Excepcion, solo si el grupo entero no tiene NINGUNA linea puramente
    numerica (comun en texto libre de Word/PDF, ej. "Paracetamol 500mg -
    120 unidades" en una sola oracion): se acepta si el numero de la
    cantidad aparece LITERALMENTE como token dentro de alguna linea citada
    — no basta con que no haya alternativa, el numero tiene que estar
    realmente ahi (si no, seria aceptar cualquier valor inventado por el
    LLM sin respaldo real, que es justo lo que este filtro existe para
    evitar)."""
    cantidad = item.get("cantidad")
    if cantidad is None:
        return True, None

    origenes = item.get("lineas_origen", []) or []
    cited_texts = [raw_lines[idx]["text"] for idx in origenes if 0 <= idx < len(raw_lines)]

    for text in cited_texts:
        m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*\$?", text.strip())
        if m and float(m.group(1)) == float(cantidad):
            return True, None

    hay_linea_numerica_en_grupo = any(
        re.fullmatch(r"\d+(?:\.\d+)?\s*\$?", l["text"].strip()) for l in raw_lines
    )
    if not hay_linea_numerica_en_grupo:
        cantidad_str = str(int(cantidad)) if float(cantidad).is_integer() else str(cantidad)
        for text in cited_texts:
            if re.search(rf"(?<!\d){re.escape(cantidad_str)}(?!\d)", text):
                return True, None

    return False, "origen_sospechoso"


def resolve_moderate_groups(rows_raw, moderate_idx):
    """Manda UNA llamada al LLM POR GRUPO moderado (<=5 lineas), sin numeros de
    grupo en el prompt (elimina la desalineacion de indice que ocurria en el
    formato batch). Delega la resolucion+verificacion de cada grupo a
    resolver_grupo_llm (compartida con el flujo de texto libre de Word/PDF)."""
    resueltos, requiere_revision = [], []
    for gi, i in enumerate(moderate_idx):
        r, rv = resolver_grupo_llm(rows_raw[i], f"grupo_{gi}")
        resueltos.extend(r)
        requiere_revision.extend(rv)
    return resueltos, requiere_revision


def lineas_a_raw_lines(lineas):
    """Convierte una lista de strings de texto plano (sin coordenadas reales,
    ej. parrafos de Word o texto extraido de un PDF) al mismo formato que
    rows_raw usa para las lineas de OCR, para poder reutilizar
    resolver_grupo_llm tal cual. "y" = indice de linea (orden de arriba hacia
    abajo), "x" = 0 (no aplica, no hay columnas reales)."""
    return [{"text": t, "x1": 0, "y1": i} for i, t in enumerate(lineas)]


def structure_free_text_with_llm(lineas, id_prefix, chunk_size=10):
    """Agente Estructurador para texto libre (sin tabla): reusa
    resolver_grupo_llm en bloques de `chunk_size` lineas (documentos largos
    se parten para no mandar un prompt gigante al LLM en una sola llamada,
    ver Prueba C del pipeline de imagenes)."""
    resueltos, requiere_revision = [], []
    for start in range(0, len(lineas), chunk_size):
        bloque = lineas[start:start + chunk_size]
        raw_lines = lineas_a_raw_lines(bloque)
        r, rv = resolver_grupo_llm(raw_lines, f"{id_prefix}_bloque{start // chunk_size}")
        resueltos.extend(r)
        requiere_revision.extend(rv)
    return resueltos, requiere_revision


# ---------------------------------------------------------------------------
# Pipeline completo
# ---------------------------------------------------------------------------

def process_image(image_path):
    t_start = time.perf_counter()

    recon = run_ocr_and_reconstruct(image_path)
    table = recon["table"]
    rows_raw = recon["rows_raw"]
    n_lines = recon["n_lines"]
    n_rows = recon["n_rows"]
    ratio = n_lines / n_rows if n_rows else 0
    global_suspicion = ratio > RATIO_THRESHOLD
    name_word_limit = 5 if global_suspicion else 7

    name_col, qty_col = pick_name_and_qty_columns(table)

    resueltos = []
    requiere_revision = []
    moderate_idx = []

    for i, row in enumerate(table):
        amb, reason = is_ambiguous(row, name_word_limit=name_word_limit)
        if not amb:
            nombre = row[name_col].strip() if name_col is not None else ""
            cantidad_raw = row[qty_col].strip() if qty_col is not None else ""
            m = re.search(r"^\s*(\d+(\.\d+)?)\s*$", cantidad_raw)
            cantidad = None
            if m:
                cantidad = m.group(1)
            else:
                # Fallback: la columna de cantidad a veces queda vacia porque el
                # numero se fusiono con la columna de precio (ej. "4 $"). Buscamos
                # ese patron en cualquier columna que no sea la de nombre.
                for c, cell in enumerate(row):
                    if c == name_col:
                        continue
                    m2 = re.match(r"^\s*(\d+(\.\d+)?)\s*\$\s*$", cell.strip())
                    if m2:
                        cantidad = m2.group(1)
                        break

            ok_nombre, _ = filtro1_nombre_valido(nombre)
            if nombre and cantidad is not None and ok_nombre:
                resueltos.append({
                    "nombre": nombre,
                    "cantidad": float(cantidad) if "." in cantidad else int(cantidad),
                    "confianza": "alta",
                })
            elif nombre and cantidad is not None and not ok_nombre:
                # Reconstruccion "clara" (sin LLM), pero el nombre en si es
                # invalido (ej. "Q" — fragmento suelto de OCR, no un
                # producto real). Antes esto se colaba directo a resueltos
                # porque el Filtro 1 solo corria sobre filas del LLM.
                requiere_revision.append({
                    "id": f"fila_{i}",
                    "texto_ocr_crudo": [it["text"] for it in rows_raw[i]],
                    "nombre_sugerido": nombre,
                    "cantidad_sugerida": cantidad,
                    "motivo": "nombre_invalido",
                })
            else:
                requiere_revision.append({
                    "id": f"fila_{i}",
                    "texto_ocr_crudo": [it["text"] for it in rows_raw[i]],
                    "nombre_sugerido": nombre or None,
                    "cantidad_sugerida": None,
                    "motivo": "cobertura_incompleta",
                })
            continue

        n_raw = len(rows_raw[i])
        if n_raw > EXTREME_FUSION_THRESHOLD:
            requiere_revision.append({
                "id": f"fila_{i}",
                "texto_ocr_crudo": [it["text"] for it in rows_raw[i]],
                "nombre_sugerido": None,
                "cantidad_sugerida": None,
                "motivo": "fusion_extrema",
            })
        else:
            moderate_idx.append(i)

    if moderate_idx:
        llm_resueltos, llm_revision = resolve_moderate_groups(rows_raw, moderate_idx)
        resueltos.extend(llm_resueltos)
        requiere_revision.extend(llm_revision)

    t_end = time.perf_counter()

    return {
        "resueltos": resueltos,
        "requiere_revision": requiere_revision,
        "tiempo_procesamiento_seg": round(t_end - t_start, 2),
    }


# ---------------------------------------------------------------------------
# Excel / Word (tablas) / PDF (tablas): texto YA estructurado en filas, sin el
# desorden del OCR -> van directo a resueltos, sin LLM ni detector de
# ambiguedad (esos existen para lidiar con fusiones de OCR que aca no
# aplican). Solo se valida que cada fila tenga una cantidad numerica real.
# ---------------------------------------------------------------------------

NOMBRE_KEYWORDS = {"nombre", "producto", "descripcion", "descripción", "articulo", "artículo", "item", "ítem", "medicamento"}
CANTIDAD_KEYWORDS = {"cantidad", "cant", "stock", "unidades", "qty", "existencias"}


def detectar_columnas_header(primera_fila):
    """Si la primera fila parece un encabezado (alguna celda coincide con
    palabras clave de nombre/cantidad), devuelve (col_nombre, col_cantidad,
    hay_header=True). Si no, asume las 2 primeras columnas (0, 1)."""
    col_nombre = col_cantidad = None
    for c, cell in enumerate(primera_fila):
        val = str(cell or "").strip().lower()
        if val in NOMBRE_KEYWORDS and col_nombre is None:
            col_nombre = c
        elif val in CANTIDAD_KEYWORDS and col_cantidad is None:
            col_cantidad = c
    if col_nombre is not None and col_cantidad is not None:
        return col_nombre, col_cantidad, True
    return 0, 1, False


def filas_a_resueltos_directos(filas, id_prefix):
    """filas: lista de filas, cada una una lista de valores de celda (ya
    convertidos a string). Ver PASO 1/2 del pipeline: Excel y tablas de
    Word/PDF comparten esta misma logica."""
    filas = [f for f in filas if any(str(c or "").strip() for c in f)]  # descarta filas totalmente vacias
    if not filas:
        return [], []

    col_nombre, col_cantidad, hay_header = detectar_columnas_header(filas[0])
    datos = filas[1:] if hay_header else filas

    resueltos, requiere_revision = [], []
    for i, fila in enumerate(datos):
        nombre = str(fila[col_nombre]).strip() if col_nombre < len(fila) and fila[col_nombre] is not None else ""
        cantidad_raw = str(fila[col_cantidad]).strip() if col_cantidad < len(fila) and fila[col_cantidad] is not None else ""

        m = re.search(r"\d+(\.\d+)?", cantidad_raw)
        cantidad = None
        if m:
            cantidad = float(m.group()) if "." in m.group() else int(m.group())

        ok_nombre, _ = filtro1_nombre_valido(nombre)
        if nombre and cantidad is not None and ok_nombre:
            resueltos.append({"nombre": nombre, "cantidad": cantidad, "confianza": "alta"})
        elif nombre and cantidad is not None and not ok_nombre:
            requiere_revision.append({
                "id": f"{id_prefix}_{i}",
                "texto_ocr_crudo": [str(c) for c in fila if str(c or "").strip()],
                "nombre_sugerido": nombre,
                "cantidad_sugerida": cantidad,
                "motivo": "nombre_invalido",
            })
        else:
            requiere_revision.append({
                "id": f"{id_prefix}_{i}",
                "texto_ocr_crudo": [str(c) for c in fila if str(c or "").strip()],
                "nombre_sugerido": nombre or None,
                "cantidad_sugerida": None,
                "motivo": "sin_cantidad",
            })
    return resueltos, requiere_revision


def procesar_excel_archivo(path):
    from openpyxl import load_workbook

    t_start = time.perf_counter()
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    filas = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    resueltos, requiere_revision = filas_a_resueltos_directos(filas, "fila")
    t_end = time.perf_counter()
    return {
        "resueltos": resueltos,
        "requiere_revision": requiere_revision,
        "tiempo_procesamiento_seg": round(t_end - t_start, 2),
    }


def procesar_word_archivo(path):
    from docx import Document

    t_start = time.perf_counter()
    doc = Document(path)
    resueltos, requiere_revision = [], []

    if doc.tables:
        # Documento con tabla(s): mismo tratamiento que Excel, una tabla a la vez.
        for ti, table in enumerate(doc.tables):
            filas = [[cell.text for cell in row.cells] for row in table.rows]
            r, rv = filas_a_resueltos_directos(filas, f"tabla{ti}_fila")
            resueltos.extend(r)
            requiere_revision.extend(rv)
    else:
        # Texto libre (parrafos, sin tabla) -> pasa por el Agente
        # Estructurador (LLM), igual que el texto ambiguo de OCR. Se excluyen
        # los parrafos con estilo "Heading"/"Title" (titulos del documento,
        # no son productos) — si no, el LLM no sabe donde meterlos y la
        # verificacion de cobertura tumba todo el bloque a revision.
        lineas = [
            p.text.strip() for p in doc.paragraphs
            if p.text.strip() and not (p.style and p.style.name or "").startswith(("Heading", "Title"))
        ]
        if lineas:
            resueltos, requiere_revision = structure_free_text_with_llm(lineas, "parrafo")

    t_end = time.perf_counter()
    return {
        "resueltos": resueltos,
        "requiere_revision": requiere_revision,
        "tiempo_procesamiento_seg": round(t_end - t_start, 2),
    }


def procesar_pdf_archivo(path):
    import pdfplumber

    t_start = time.perf_counter()
    resueltos, requiere_revision = [], []

    with pdfplumber.open(path) as pdf:
        tablas_encontradas = []
        texto_por_pagina = []
        for page in pdf.pages:
            tablas_encontradas.extend(page.extract_tables() or [])
            texto_por_pagina.append(page.extract_text() or "")
        texto_total = "\n".join(texto_por_pagina).strip()
        n_paginas = len(pdf.pages)

    if tablas_encontradas:
        # PDF con tabla(s) detectable(s) (ej. exportado desde Excel/Word) ->
        # mismo tratamiento directo que Excel.
        for ti, tabla in enumerate(tablas_encontradas):
            r, rv = filas_a_resueltos_directos(tabla, f"tabla{ti}_fila")
            resueltos.extend(r)
            requiere_revision.extend(rv)

    elif len(texto_total) > 20:
        # Tiene texto seleccionable real (no es una imagen escaneada) pero sin
        # estructura de tabla -> Agente Estructurador, igual que Word libre.
        lineas = [l.strip() for l in texto_total.splitlines() if l.strip()]
        resueltos, requiere_revision = structure_free_text_with_llm(lineas, "linea")

    else:
        # // TODO: no probado con un PDF escaneado real todavia (no habia
        # uno a mano) — este camino no deberia romper nada si nunca se
        # activa, pero tampoco tiene una prueba end-to-end que lo confirme.
        # Sin texto extraible -> es una imagen escaneada metida en el PDF.
        # Convertimos cada pagina a imagen (PyMuPDF, sin depender de Poppler)
        # y reusamos el pipeline de OCR de imagenes tal cual (process_image).
        import fitz  # PyMuPDF

        pdf_doc = fitz.open(path)
        for pi, page in enumerate(pdf_doc):
            pix = page.get_pixmap(dpi=200)
            tmp_img = path.with_name(f"{path.stem}_pagina{pi}_{uuid.uuid4().hex}.jpg")
            pix.save(str(tmp_img))
            try:
                res_pagina = process_image(tmp_img)
                resueltos.extend(res_pagina["resueltos"])
                requiere_revision.extend(res_pagina["requiere_revision"])
            finally:
                if tmp_img.exists():
                    tmp_img.unlink()
        pdf_doc.close()

    t_end = time.perf_counter()
    return {
        "resueltos": resueltos,
        "requiere_revision": requiere_revision,
        "tiempo_procesamiento_seg": round(t_end - t_start, 2),
    }


# ---------------------------------------------------------------------------
# Agente Supervisor: clasificacion de intencion + respuesta a preguntas del
# sistema (ver scripts/conocimiento_sistema.md). Solo se llama cuando el
# mensaje NO trae archivos adjuntos (con archivos, es casi seguro
# "importar" y ni vale la pena preguntarle al LLM).
# ---------------------------------------------------------------------------

CONOCIMIENTO_PATH = Path(__file__).parent / "conocimiento_sistema.md"

CATEGORIAS_INTENCION = ("importar", "pregunta_sistema", "charla_general")

CHARLA_GENERAL_RESPUESTA = (
    "¡Hola! Puedo ayudarte a cargar tu inventario a partir de fotos, Excel, Word o PDF "
    "— solo adjuntá el archivo con el botón \"+\" de abajo. También puedo responder "
    "preguntas sobre cómo usar el sistema."
)

IMPORTAR_SIN_ARCHIVO_RESPUESTA = (
    "Para importar datos, adjuntá una foto o un archivo (Excel, Word o PDF) de tu "
    "inventario con el botón \"+\" de abajo, y yo me encargo del resto."
)


def cargar_conocimiento():
    try:
        return CONOCIMIENTO_PATH.read_text(encoding="utf-8")
    except FileNotFoundError:
        return ""


def clasificar_intencion(texto):
    prompt = f"""Clasifica la intencion de este mensaje de un usuario, escrito en el chat
de un sistema de gestion de gimnasio, en EXACTAMENTE una de estas 3 categorias:

- "importar": quiere cargar/importar/subir datos de inventario (aunque no haya adjuntado
  ningun archivo todavia, ej. "quiero cargar mi inventario", "como subo una lista de productos").
- "pregunta_sistema": esta preguntando como usar o donde encontrar algo del sistema (ej.
  clientes, ventas, caja, membresias, inventario, casilleros, reportes, usuarios).
- "charla_general": saludo, agradecimiento, o cualquier charla que no tiene que ver con
  importar datos ni con como usar el sistema.

Mensaje del usuario: "{texto}"

Responde SOLO con una palabra, exactamente una de: importar, pregunta_sistema,
charla_general. Nada mas, sin explicacion.
"""
    result = call_ollama(prompt)
    raw = result.get("response", "").strip().lower()
    for cat in CATEGORIAS_INTENCION:
        if cat in raw:
            return cat
    return "charla_general"


# Preguntas por datos externos en tiempo real (fecha, hora, clima, noticias) que un LLM
# chico como qwen2.5:3b tiende a "inventar" en vez de admitir que no tiene acceso a eso.
# Se detectan por palabra clave y se responden de forma honesta y determinista, sin pasar
# por el LLM, en vez de confiar en que el modelo siga esa instruccion de forma consistente.
PATRON_DATO_EXTERNO = re.compile(
    r"\bque (dia|hora|fecha)\b|\bque hora es\b|\bclima\b|\btemperatura afuera\b|\bnoticias?\b|\bpronostico del tiempo\b",
    re.IGNORECASE,
)

RESPUESTA_DATO_EXTERNO = (
    "No tengo acceso a esa información en tiempo real (fecha, hora, clima, etc.). "
    "Puedo ayudarte a cargar tu inventario a partir de fotos, Excel, Word o PDF, o "
    "responder preguntas sobre cómo usar el sistema."
)


def responder_charla_general(texto):
    if PATRON_DATO_EXTERNO.search(texto):
        return RESPUESTA_DATO_EXTERNO

    prompt = f"""Eres el asistente de un sistema de gestion de gimnasio, en el chat de
importacion de inventario. El usuario escribio un saludo, agradecimiento, o charla simple
que no tiene que ver con importar datos ni con como usar el sistema.

Respondele en espanol, en 1 a 3 frases, con tono simple y amable, y mencion que podes
ayudar a cargar su inventario (fotos, Excel, Word o PDF) o responder preguntas sobre el
sistema. No inventes funciones del sistema que no existan.

Mensaje del usuario: "{texto}"
"""
    result = call_ollama(prompt)
    respuesta = result.get("response", "").strip()
    return respuesta or CHARLA_GENERAL_RESPUESTA


def responder_pregunta_sistema(texto):
    conocimiento = cargar_conocimiento()
    prompt = f"""Eres el asistente de un sistema de gestion de gimnasio. Responde la
pregunta del usuario usando SOLO la informacion del documento de abajo — no inventes
funciones ni pasos que no esten ahi. Si la pregunta no esta cubierta en el documento, decilo
honestamente (ej. "todavia no tengo esa informacion") en vez de adivinar una respuesta.
Responde en espanol, en 2 a 4 frases, tono simple y amable.

--- DOCUMENTO ---
{conocimiento}
--- FIN DOCUMENTO ---

Pregunta del usuario: {texto}
"""
    result = call_ollama(prompt)
    return result.get("response", "").strip()


def procesar_asistente_texto(texto):
    intencion = clasificar_intencion(texto)
    if intencion == "pregunta_sistema":
        respuesta = responder_pregunta_sistema(texto)
    elif intencion == "importar":
        respuesta = IMPORTAR_SIN_ARCHIVO_RESPUESTA
    else:
        respuesta = responder_charla_general(texto)
    return {"intencion": intencion, "respuesta": respuesta}


# ---------------------------------------------------------------------------
# Endpoints HTTP
# ---------------------------------------------------------------------------

@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "ok", "modelo_llm": OLLAMA_MODEL})


@app.route("/procesar-imagen", methods=["POST"])
def procesar_imagen():
    tmp_path = None
    try:
        if "imagen" in request.files:
            f = request.files["imagen"]
            tmp_path = Path(f"./_tmp_upload_{uuid.uuid4().hex}.jpg")
            f.save(tmp_path)
        else:
            data = request.get_json(silent=True) or {}
            b64 = data.get("imagen_base64")
            if not b64:
                return jsonify({"error": "falta 'imagen' (multipart) o 'imagen_base64' (json)"}), 400
            if "," in b64[:50]:
                b64 = b64.split(",", 1)[1]
            raw = base64.b64decode(b64)
            tmp_path = Path(f"./_tmp_upload_{uuid.uuid4().hex}.jpg")
            tmp_path.write_bytes(raw)

        result = process_image(tmp_path)
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if tmp_path and tmp_path.exists():
            tmp_path.unlink()


def _guardar_archivo_subido(campo, extension):
    """Comun a los 3 endpoints nuevos: guarda el archivo subido (multipart,
    campo `campo`) a un temporal y devuelve su Path. Lanza ValueError si no
    vino ningun archivo (se traduce a 400 en el endpoint)."""
    if campo not in request.files:
        raise ValueError(f"falta el campo '{campo}' (multipart/form-data)")
    f = request.files[campo]
    tmp_path = Path(f"./_tmp_upload_{uuid.uuid4().hex}{extension}")
    f.save(tmp_path)
    return tmp_path


@app.route("/procesar-excel", methods=["OPTIONS"])
@app.route("/procesar-word", methods=["OPTIONS"])
@app.route("/procesar-pdf", methods=["OPTIONS"])
def _preflight_otros_formatos():
    return ("", 204)


@app.route("/procesar-excel", methods=["POST"])
def procesar_excel():
    tmp_path = None
    try:
        tmp_path = _guardar_archivo_subido("archivo", ".xlsx")
        return jsonify(procesar_excel_archivo(tmp_path))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if tmp_path and tmp_path.exists():
            tmp_path.unlink()


@app.route("/procesar-word", methods=["POST"])
def procesar_word():
    tmp_path = None
    try:
        tmp_path = _guardar_archivo_subido("archivo", ".docx")
        return jsonify(procesar_word_archivo(tmp_path))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if tmp_path and tmp_path.exists():
            tmp_path.unlink()


@app.route("/procesar-pdf", methods=["POST"])
def procesar_pdf():
    tmp_path = None
    try:
        tmp_path = _guardar_archivo_subido("archivo", ".pdf")
        return jsonify(procesar_pdf_archivo(tmp_path))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if tmp_path and tmp_path.exists():
            tmp_path.unlink()


@app.route("/asistente-texto", methods=["OPTIONS"])
def _asistente_texto_preflight():
    # JSON (a diferencia de multipart/form-data) SI dispara un preflight CORS
    # real en el navegador — sin esta ruta, Chromium bloquea la llamada
    # antes de que llegue al POST de abajo.
    return ("", 204)


@app.route("/asistente-texto", methods=["POST"])
def asistente_texto():
    try:
        data = request.get_json(silent=True) or {}
        texto = (data.get("texto") or "").strip()
        if not texto:
            return jsonify({"error": "falta 'texto'"}), 400
        return jsonify(procesar_asistente_texto(texto))
    except Exception as e:
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    print(f"Servicio OCR escuchando en http://localhost:{PORT}")
    app.run(host="127.0.0.1", port=PORT, debug=False)
