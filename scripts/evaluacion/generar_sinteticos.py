"""Genera 39 documentos sintéticos en español con respuesta conocida (Sprint 2, paso 1.2).

Uso:  venv-ocr\\Scripts\\python.exe scripts\\evaluacion\\generar_sinteticos.py [--salida datos-evaluacion/sinteticos]

Cada documento <id>.<ext> va con <id>.json:
  { "id", "tipo", "archivo", "items": [{nombre, cantidad, unidad?, revision_esperada?, motivo?, duplicado?}],
    "no_inventario": [textos que NO deben salir como producto] }
Todo es determinista (semilla fija): volver a correrlo da los mismos archivos.
Tipos: lista_simple(6) tabla_limpia(6) tabla_compleja(6) manuscrito(6) degradada(6) excel(3) word(3) pdf(3).
"""
import argparse
import json
import random
import sys
import textwrap
from pathlib import Path

import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont

FUENTES = Path("C:/Windows/Fonts")
rng = random.Random(20261030)


def fuente(nombre, tam):
    return ImageFont.truetype(str(FUENTES / nombre), tam)


# ── Catálogos ────────────────────────────────────────────────────────────────────────────────────────────────
GYM = ["Mancuerna hexagonal 5 kg", "Proteína whey 2 lb", "Creatina monohidratada 300 g", "Banda elástica de resistencia",
       "Guantes de entrenamiento", "Cinturón lumbar", "Toalla deportiva", "Shaker 700 ml", "Cuerda para saltar",
       "Colchoneta de yoga 6 mm", "Disco olímpico 10 kg", "Agua mineral 500 ml", "Bebida isotónica 600 ml",
       "Barra de proteína", "Rodillera ajustable", "Muñequeras", "Candado para casillero", "Tobilleras con peso 2 kg",
       "Pre-entreno 30 servicios", "Camiseta dry-fit talla M", "Botella térmica 1 litro", "Magnesio en cápsulas x60"]
FARMACIA = ["Paracetamol 500mg", "Ibuprofeno 400mg", "Amoxicilina 500mg x20", "Omeprazol 20mg", "Loratadina 10mg",
            "Alcohol 70° 500 ml", "Algodón 100 g", "Gasas estériles x10", "Jarabe para la tos 120 ml", "Vitamina C 1g x30",
            "Suero oral", "Termómetro digital", "Jeringa 5 ml", "Diclofenaco gel 60 g", "Metformina 850mg", "Enalapril 10mg",
            "Cetirizina 10mg", "Vendas elásticas 10 cm", "Mascarilla quirúrgica x50", "Ácido fólico 5mg", "Ampolla de B12 x3",
            "Crema para niños 100 g"]
ABARROTES = ["Coca cola light de 12 pzas 355 ml", "Arroz extra", "Azúcar blanca", "Aceite vegetal", "Leche entera",
             "Galletas de soda", "Atún en agua", "Fideos spaghetti", "Papel higiénico x12", "Detergente en polvo",
             "Jabón de tocador", "Pañales desechables premium", "Servilletas", "Sal yodada", "Harina de trigo", "Mantequilla",
             "Yogur natural", "Cereal de maíz", "Salsa de tomate", "Mayonesa", "Café molido", "Té de manzanilla",
             "Cerveza lata 355 ml", "Jugo de naranja", "Huevos x30", "Piña en almíbar", "Cañahua tostada", "Quinua real"]
ABREV = [("Prot. whey 2lb", 14), ("Cr. monohidrat. 300g", 9), ("Amox. 500mg", 40), ("Ibup. 400mg", 55), ("Det. polvo 1kg", 12)]


def tomar(pool, n):
    return rng.sample(pool, n)


def cant():
    return rng.randint(2, 96)


def it(nombre, cantidad=None, unidad=None, **extra):
    d = {"nombre": nombre, "cantidad": cantidad if cantidad is not None else cant()}
    if unidad:
        d["unidad"] = unidad
    d.update(extra)
    return d


def marcar_decimal(item):
    item["revision_esperada"] = True
    item["motivo"] = "decimal" if isinstance(item["cantidad"], float) else "unidad"
    return item


# ── Dibujo de listas y tablas en imagen ──────────────────────────────────────────────────────────────────────
def lienzo(w, h, fondo=(255, 255, 255)):
    im = Image.new("RGB", (w, h), fondo)
    return im, ImageDraw.Draw(im)


def texto_cant(item):
    c = item["cantidad"]
    s = f"{c:g}" if isinstance(c, float) else str(c)
    return f"{s} {item['unidad']}" if item.get("unidad") else s


def render_lista(filas, titulo, fnt="arial.ttf", tam=34, ancho=1240, color=(15, 15, 15), fondo=(255, 255, 255), cabecera=None):
    """Lista impresa de dos columnas sin bordes: nombre a la izquierda, cantidad a la derecha.
    filas: [{kind: item|total|vacia|nota, item?, texto?}]. Devuelve (imagen, no_inventario)."""
    f, fb = fuente(fnt, tam), fuente(fnt.replace(".ttf", "bd.ttf") if "arial" in fnt or "times" in fnt else fnt, tam + 6)
    paso = int(tam * 1.9)
    alto = 150 + paso * (len(filas) + 2) + 80
    im, d = lienzo(ancho, alto, fondo)
    no_inv = []
    d.text((70, 50), titulo, font=fb, fill=color)
    no_inv.append(titulo)
    y = 130
    if cabecera:
        d.text((70, y), cabecera[0], font=f, fill=color)
        d.text((ancho - 380, y), cabecera[1], font=f, fill=color)
        no_inv.extend(cabecera)
        y += paso
    for fila in filas:
        k = fila["kind"]
        if k == "item":
            d.text((70, y), fila["item"]["nombre"], font=f, fill=color)
            d.text((ancho - 380, y), texto_cant(fila["item"]), font=f, fill=color)
        elif k == "total":
            d.line([(70, y - 6), (ancho - 120, y - 6)], fill=color, width=2)
            d.text((70, y), fila["texto"], font=fb, fill=color)
            no_inv.append(fila["texto"])
        elif k == "nota":
            d.text((70, y), fila["texto"], font=f, fill=color)
            no_inv.append(fila["texto"])
        y += paso
    return im, no_inv


def render_tabla(columnas, filas, titulo=None, fnt="arial.ttf", tam=28, anchos=None, ancho=1400, color=(15, 15, 15), fondo=(255, 255, 255),
                 grosor=2, cab_fondo=(225, 228, 235)):
    """Tabla con bordes. filas: [{kind: item|categoria|total|vacia, celdas: [texto por columna] (las celdas pueden ser multilínea),
    item?}]. Las celdas se parten en varias líneas según el ancho de columna. Devuelve (imagen, no_inventario)."""
    f = fuente(fnt, tam)
    fb = fuente("arialbd.ttf" if "arial" in fnt else fnt, tam)
    anchos = anchos or [int((ancho - 80) / len(columnas))] * len(columnas)
    x0 = (ancho - sum(anchos)) // 2
    interlin = int(tam * 1.25)
    pad = 12

    def partir(txt, w):
        out = []
        for parr in str(txt).split("\n"):
            palabras, linea = parr.split(" "), ""
            for p in palabras:
                prueba = (linea + " " + p).strip()
                if d0.textlength(prueba, font=f) <= w - 2 * pad:
                    linea = prueba
                else:
                    if linea:
                        out.append(linea)
                    linea = p
            out.append(linea)
        return out

    d0 = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    celdas = [("cab", columnas)] + [(fila["kind"], fila["celdas"]) for fila in filas]
    alturas = []
    for kind, cs in celdas:
        if kind == "categoria":
            alturas.append(interlin + 2 * pad)
        else:
            alturas.append(max(len(partir(c, anchos[i])) for i, c in enumerate(cs)) * interlin + 2 * pad)
    top = 140 if titulo else 60
    alto = top + sum(alturas) + 100
    im, d = lienzo(ancho, alto, fondo)
    no_inv = []
    if titulo:
        d.text((x0, 55), titulo, font=fuente("arialbd.ttf", tam + 8), fill=color)
        no_inv.append(titulo)
    y = top
    for (kind, cs), h in zip(celdas, alturas):
        if kind == "cab":
            d.rectangle([x0, y, x0 + sum(anchos), y + h], fill=cab_fondo)
            no_inv.extend([c for c in columnas if c])
        x = x0
        if kind == "categoria":
            d.rectangle([x0, y, x0 + sum(anchos), y + h], outline=color, width=grosor, fill=(240, 240, 240))
            d.text((x0 + pad, y + pad), cs[0], font=fb, fill=color)
            no_inv.append(cs[0])
        else:
            for i, c in enumerate(cs):
                d.rectangle([x, y, x + anchos[i], y + h], outline=color, width=grosor)
                for j, ln in enumerate(partir(c, anchos[i])):
                    d.text((x + pad, y + pad + j * interlin), ln, font=fb if kind in ("cab", "total") else f, fill=color)
                x += anchos[i]
            if kind == "total":
                no_inv.extend([c for c in cs if c.strip()])
        y += h
    return im, no_inv


def render_manuscrito(filas, titulo, fnt, tam=46, ancho=1240, tinta=(25, 35, 90)):
    f = fuente(fnt, tam)
    paso = int(tam * 1.7)
    alto = 160 + paso * (len(filas) + 1)
    im, d = lienzo(ancho, alto, (250, 247, 232))
    for y in range(120, alto, paso):  # renglones de cuaderno
        d.line([(40, y + paso - 8), (ancho - 40, y + paso - 8)], fill=(185, 205, 225), width=2)
    d.line([(110, 0), (110, alto)], fill=(230, 160, 160), width=3)
    d.text((140, 40), titulo, font=f, fill=tinta)
    y = 120
    for fila in filas:
        j = rng.randint(-3, 3)
        d.text((140, y + j), fila["item"]["nombre"], font=f, fill=tinta)
        d.text((ancho - 330, y + rng.randint(-3, 3)), texto_cant(fila["item"]), font=f, fill=tinta)
        y += paso
    return im, [titulo]


# ── Degradaciones ────────────────────────────────────────────────────────────────────────────────────────────
def a_cv(im):
    return cv2.cvtColor(np.array(im), cv2.COLOR_RGB2BGR)


def a_pil(a):
    return Image.fromarray(cv2.cvtColor(a, cv2.COLOR_BGR2RGB))


def degradar(im, modo):
    a = a_cv(im)
    h, w = a.shape[:2]
    if modo == "rotacion":
        M = cv2.getRotationMatrix2D((w / 2, h / 2), 6.5, 1.0)
        a = cv2.warpAffine(a, M, (w, h), borderValue=(235, 235, 235))
    elif modo == "perspectiva":
        src = np.float32([[0, 0], [w, 0], [w, h], [0, h]])
        dst = np.float32([[60, 30], [w - 20, 0], [w - 70, h - 25], [10, h - 10]])
        a = cv2.warpPerspective(a, cv2.getPerspectiveTransform(src, dst), (w, h), borderValue=(215, 215, 215))
    elif modo == "sombra":
        yy, xx = np.mgrid[0:h, 0:w]
        sombra = np.clip(0.45 + 0.55 * (xx / w) ** 0.8, 0, 1)  # gradiente fuerte: un lado casi a la mitad de luz
        a = (a.astype(np.float32) * sombra[..., None]).astype(np.uint8)
    elif modo == "desenfoque":
        a = cv2.GaussianBlur(a, (0, 0), 2.6)
    elif modo == "jpeg_bajo":
        a = cv2.resize(a, (w // 2, h // 2), interpolation=cv2.INTER_AREA)
        _, buf = cv2.imencode(".jpg", a, [cv2.IMWRITE_JPEG_QUALITY, 7])
        a = cv2.resize(cv2.imdecode(buf, 1), (w, h))
    elif modo == "bajo_contraste":
        a = (128 + (a.astype(np.float32) - 128) * 0.22 + 70).clip(0, 255).astype(np.uint8)
    return a_pil(a)


# ── Especificaciones ─────────────────────────────────────────────────────────────────────────────────────────
def filas_items(items, extra_total=None, vacias=False):
    filas = []
    for k, i in enumerate(items):
        filas.append({"kind": "item", "item": i})
        if vacias and k in (2, 5):
            filas.append({"kind": "vacia"})
    if extra_total:
        filas.append({"kind": "total", "texto": extra_total})
    return filas


def especificar():
    docs = []

    def add(tipo, **kw):
        docs.append({"tipo": tipo, **kw})

    # ---- listas simples impresas (6) ----
    a = [it(n) for n in tomar(GYM, 8)]
    add("lista_simple", titulo="INVENTARIO GIMNASIO", items=a, fnt="arial.ttf")
    b = [it(n) for n in tomar(FARMACIA, 8)]
    add("lista_simple", titulo="Inventario de farmacia", items=b, fnt="times.ttf", total="TOTAL")
    c = [it("Pañuelos desechables"), it("Niño sano: crema 50 g"), it("Sábana de camilla"), it("Jeringa de insulina"), it("Piñón de cadena")] + [it(n) for n in tomar(GYM, 3)]
    add("lista_simple", titulo="Almacén — compras de enero", items=c, fnt="calibri.ttf")
    d_ = [it("Paracetamol 500mg", 120), it("Ibuprofeno 400mg", 80), it("Amoxicilina 500mg x20", 35), it("Omeprazol 20mg", 60), it("Vitamina C 1g x30", 24), it("Gasas estériles x10", 48)]
    add("lista_simple", titulo="FARMACIA — STOCK", items=d_, fnt="consola.ttf", tam=30)
    e = [it(n) for n in tomar(GYM, 6)]
    e.insert(3, it(e[1]["nombre"], 7, duplicado=True))
    e[1]["duplicado"] = True
    add("lista_simple", titulo="Inventario gimnasio (con repetido)", items=e, fnt="verdana.ttf", tam=30)
    f_ = [it("Arroz extra", 2.5, "kg"), it("Aceite vegetal", 3, "litros"), it("Harina de trigo", 1.5, "kg"), it("Leche entera", 24, "litros"), it("Azúcar blanca", 18)]
    marcar_decimal(f_[0]); marcar_decimal(f_[1]); marcar_decimal(f_[2]); marcar_decimal(f_[3])
    add("lista_simple", titulo="PEDIDO SEMANAL", items=f_, fnt="tahoma.ttf", total="TOTAL: 5 artículos", vacias=True)

    # ---- tablas limpias (6): nombre / cantidad / precio ----
    pm = ["$", "$", "$", "Bs", "Bs", "Bs"]
    pools = [GYM, FARMACIA, ABARROTES, GYM, FARMACIA, ABARROTES]
    for k in range(6):
        items = [it(n) for n in tomar(pools[k], 8)]
        add("tabla_limpia", titulo=["Existencias gimnasio", "Botiquín y farmacia", "Productos de tienda", "Inventario deportivo", "Medicamentos", "Abarrotes"][k],
            items=items, moneda=pm[k], total=(k % 3 == 2))

    # ---- tablas con celdas multilínea, categorías y columnas que no son inventario (réplica "tienda de abarrotes") ----
    cats = [("LÁCTEOS Y HUEVOS", ["Leche entera", "Yogur natural", "Mantequilla", "Huevos x30"]),
            ("BEBIDAS", ["Coca cola light de 12 pzas 355 ml", "Jugo de naranja", "Cerveza lata 355 ml"]),
            ("LIMPIEZA", ["Detergente en polvo", "Papel higiénico x12", "Pañales desechables premium"]),
            ("DESPENSA", ["Arroz extra", "Azúcar blanca", "Fideos spaghetti", "Atún en agua"]),
            ("FARMACIA", ["Paracetamol 500mg", "Amoxicilina 500mg x20", "Jarabe para la tos 120 ml"]),
            ("SUPLEMENTOS", ["Proteína whey 2 lb", "Creatina monohidratada 300 g", "Pre-entreno 30 servicios"])]
    obs_pool = ["revisar fecha de vencimiento", "reponer pronto", "caja dañada", "pedido pendiente al proveedor", "OK", "faltan 2 en bodega"]
    for k in range(6):
        c1, c2 = cats[k], cats[(k + 3) % 6]
        items = []
        grupos = []
        for nombre_cat, nombres in (c1, c2):
            lst = [it(n) for n in nombres]
            items.extend(lst)
            grupos.append((nombre_cat, lst))
        if k == 1:
            items.append(it("Arroz extra", 2.5, "kg")); marcar_decimal(items[-1]); grupos[-1][1].append(items[-1])
        if k == 4:
            d2 = it(items[0]["nombre"], 5, duplicado=True); items[0]["duplicado"] = True; items.append(d2); grupos[-1][1].append(d2)
        add("tabla_compleja", titulo="Control de tienda " + ["enero", "febrero", "marzo", "abril", "mayo", "junio"][k], items=items, grupos=grupos,
            obs=[rng.choice(obs_pool) for _ in items], fechas=[f"{rng.randint(1, 28):02d}/{rng.randint(1, 12):02d}/2026" for _ in items])

    # ---- manuscritos simulados (6) ----
    fm = ["Inkfree.ttf", "segoesc.ttf", "segoepr.ttf", "BRADHITC.TTF", "LHANDW.TTF", "comic.ttf"]
    nm = [["Mancuerna 5 kg", "Toalla", "Shaker", "Cuerda", "Guantes", "Colchoneta"], ["Paracetamol", "Ibuprofeno", "Alcohol", "Algodón", "Suero oral", "Jeringas"],
          ["Arroz", "Azúcar", "Aceite", "Fideos", "Sal", "Harina", "Leche"], ["Proteína", "Creatina", "Barras", "Agua", "Magnesio", "Camisetas"],
          ["Café", "Té", "Galletas", "Atún", "Jabón", "Cañahua"], ["Rodilleras", "Muñequeras", "Candados", "Bandas", "Discos", "Barras"]]
    for k in range(6):
        add("manuscrito", titulo=["Inventario", "Farmacia", "Tienda", "Gimnasio", "Pedido", "Equipos"][k], items=[it(n, rng.randint(2, 60)) for n in nm[k]], fnt=fm[k])

    # ---- fotos degradadas (6): base impresa + defecto ----
    modos = ["rotacion", "perspectiva", "sombra", "desenfoque", "jpeg_bajo", "bajo_contraste"]
    for k, m in enumerate(modos):
        pool = [GYM, FARMACIA, ABARROTES][k % 3]
        add("degradada", titulo="Inventario " + m.replace("_", " "), items=[it(n) for n in tomar(pool, 7)], modo=m, base="lista" if k % 2 == 0 else "tabla", moneda="$")

    # ---- archivos (9) ----
    x1 = [it(n) for n in tomar(GYM, 8)]
    add("excel", variante=1, items=x1)
    x2_items = [it(n) for n in tomar(ABARROTES, 6)] + [it(n) for n in tomar(FARMACIA, 3)]
    add("excel", variante=2, items=x2_items, titulo="Control general")
    x3 = [it("Arroz extra", 2.5, "kg"), it("Aceite vegetal", 3, "litros"), it("Café molido", 12), it("Té de manzanilla", 30), it("Cañahua tostada", 8), it("Prot. whey 2lb", 14), it("Ibup. 400mg", 55)]
    for x in x3[:2]:
        marcar_decimal(x)
    x3.append(it("Café molido", 5, duplicado=True)); x3[2]["duplicado"] = True
    add("excel", variante=3, items=x3)
    add("word", variante=1, items=[it(n) for n in tomar(FARMACIA, 7)], titulo="Listado de medicamentos")
    add("word", variante=2, items=[it("Paracetamol 500mg", 120), it("Ibuprofeno 400mg", 80), it("Amoxicilina 500mg x20", 35), it("Ácido fólico 5mg", 42), it("Alcohol 70° 500 ml", 15)], titulo="Pedido a proveedor")
    add("word", variante=3, items=[it(n) for n in tomar(GYM, 6)], titulo="Inventario de gimnasio")
    add("pdf", variante="texto", items=[it(n) for n in tomar(GYM, 8)], titulo="Inventario de gimnasio")
    add("pdf", variante="tabla", items=[it(n) for n in tomar(ABARROTES, 8)], titulo="Existencias de la tienda")
    add("pdf", variante="escaneado", items=[it(n) for n in tomar(FARMACIA, 7)], titulo="Farmacia escaneada")
    return docs


# ── Escritura de cada documento ──────────────────────────────────────────────────────────────────────────────
def guardar_img(im, ruta, calidad=92):
    ruta = Path(ruta)
    if ruta.suffix.lower() in (".jpg", ".jpeg"):
        im.save(ruta, quality=calidad)
    else:
        im.save(ruta)


def generar(spec, ident, carpeta):
    tipo = spec["tipo"]
    items = spec["items"]
    no_inv = []
    ext = ".png"
    if tipo == "lista_simple":
        filas = filas_items(items, spec.get("total"), spec.get("vacias"))
        im, no_inv = render_lista(filas, spec["titulo"], fnt=spec["fnt"], tam=spec.get("tam", 34))
        guardar_img(im, carpeta / (ident + ext))
    elif tipo == "tabla_limpia":
        m = spec["moneda"]
        filas = [{"kind": "item", "item": i, "celdas": [i["nombre"], str(i["cantidad"]), f"{m} {rng.randint(3, 180)}.{rng.choice(['00', '50', '90'])}"]} for i in items]
        if spec.get("total"):
            filas.append({"kind": "total", "celdas": ["TOTAL", str(sum(i["cantidad"] for i in items)), ""]})
        im, no_inv = render_tabla(["Producto", "Cantidad", "Precio"], filas, titulo=spec["titulo"], anchos=[640, 260, 260])
        guardar_img(im, carpeta / (ident + ext))
    elif tipo == "tabla_compleja":
        filas, k = [], 0
        for cat, lst in spec["grupos"]:
            filas.append({"kind": "categoria", "celdas": [cat, "", "", "", ""]})
            for i in lst:
                filas.append({"kind": "item", "item": i, "celdas": [i["nombre"], texto_cant(i), f"$ {rng.randint(3, 99)}.00", spec["fechas"][k], spec["obs"][k]]})
                k += 1
        filas.append({"kind": "total", "celdas": ["TOTAL", str(sum(int(i["cantidad"]) for i in items)), "", "", ""]})
        im, no_inv = render_tabla(["Descripción", "Cantidad", "Precio", "Fecha de ingreso", "Observaciones"], filas, titulo=spec["titulo"], anchos=[330, 170, 150, 200, 360], ancho=1400)
        no_inv += spec["obs"] + spec["fechas"]
        guardar_img(im, carpeta / (ident + ext))
    elif tipo == "manuscrito":
        im, no_inv = render_manuscrito([{"item": i} for i in items], spec["titulo"], spec["fnt"])
        guardar_img(im, carpeta / (ident + ext))
    elif tipo == "degradada":
        if spec["base"] == "lista":
            im, no_inv = render_lista(filas_items(items), spec["titulo"], tam=36)
        else:
            m = spec["moneda"]
            filas = [{"kind": "item", "item": i, "celdas": [i["nombre"], str(i["cantidad"]), f"{m} {rng.randint(3, 180)}.00"]} for i in items]
            im, no_inv = render_tabla(["Producto", "Cantidad", "Precio"], filas, titulo=spec["titulo"], anchos=[640, 260, 260])
        im = degradar(im, spec["modo"])
        ext = ".jpg"
        guardar_img(im, carpeta / (ident + ext), calidad=85)
    elif tipo == "excel":
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        v = spec["variante"]
        if v == 1:
            ws.append(["Producto", "Cantidad"]); no_inv = ["Producto", "Cantidad"]
            for i in items:
                ws.append([i["nombre"], i["cantidad"]])
        elif v == 2:
            ws.append([spec["titulo"]]); ws.append([])
            ws.append(["Descripción", "Cantidad", "Fecha", "Observaciones"]); no_inv = [spec["titulo"], "Descripción", "Cantidad", "Fecha", "Observaciones"]
            for n, i in enumerate(items):
                if n in (0, 6):
                    ws.append(["ABARROTES" if n == 0 else "FARMACIA"]); no_inv.append("ABARROTES" if n == 0 else "FARMACIA")
                ob = rng.choice(obs_pool := ["revisar vencimiento", "reponer", "OK"]); ws.append([i["nombre"], i["cantidad"], f"{rng.randint(1, 28):02d}/03/2026", ob]); no_inv.append(ob)
            ws.append([]); ws.append(["TOTAL", sum(i["cantidad"] for i in items)]); no_inv.append("TOTAL")
        else:
            ws.append(["Producto", "Existencias", "Unidad"]); no_inv = ["Producto", "Existencias", "Unidad"]
            for i in items:
                ws.append([i["nombre"], i["cantidad"], i.get("unidad", "pza")])
            ws.append([None, None, None]); ws.append(["TOTAL", None, None]); no_inv.append("TOTAL")
        ext = ".xlsx"; wb.save(carpeta / (ident + ext))
    elif tipo == "word":
        from docx import Document
        doc = Document(); v = spec["variante"]
        doc.add_heading(spec["titulo"], level=1); no_inv = [spec["titulo"]]
        if v == 1 or v == 3:
            if v == 3:
                doc.add_paragraph("Revisado por la administración el lunes. Los datos corresponden al cierre del mes.")
                no_inv.append("Revisado por la administración el lunes. Los datos corresponden al cierre del mes.")
            t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"
            t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Producto", "Cantidad"; no_inv += ["Producto", "Cantidad"]
            for i in items:
                r = t.add_row().cells; r[0].text, r[1].text = i["nombre"], str(i["cantidad"])
            if v == 3:
                doc.add_paragraph("Observación: reponer lo antes posible."); no_inv.append("Observación: reponer lo antes posible.")
        else:
            doc.add_paragraph("Detalle del pedido:"); no_inv.append("Detalle del pedido:")
            for i in items:
                doc.add_paragraph(f"{i['nombre']} - {i['cantidad']} unidades")
        ext = ".docx"; doc.save(carpeta / (ident + ext))
    elif tipo == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        pdfmetrics.registerFont(TTFont("Arial", str(FUENTES / "arial.ttf")))
        v = spec["variante"]; ruta = carpeta / (ident + ".pdf"); no_inv = [spec["titulo"]]
        if v == "texto":
            c = canvas.Canvas(str(ruta), pagesize=A4); c.setFont("Arial", 16); c.drawString(60, 790, spec["titulo"])
            c.setFont("Arial", 12); y = 750
            for i in items:
                c.drawString(60, y, f"{i['nombre']}  -  {i['cantidad']} unidades"); y -= 26
            c.save()
        elif v == "tabla":
            data = [["Producto", "Cantidad"]] + [[i["nombre"], str(i["cantidad"])] for i in items]; no_inv += ["Producto", "Cantidad"]
            sd = SimpleDocTemplate(str(ruta), pagesize=A4)
            tabla = Table(data, colWidths=[330, 100]); tabla.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.7, colors.black), ("FONTNAME", (0, 0), (-1, -1), "Arial")]))
            sd.build([Paragraph(spec["titulo"], getSampleStyleSheet()["Title"]), Spacer(1, 12), tabla])
        else:  # escaneado: página-imagen sin texto seleccionable
            filas = [{"kind": "item", "item": i, "celdas": [i["nombre"], str(i["cantidad"])]} for i in items]
            im, no_inv = render_tabla(["Producto", "Cantidad"], filas, titulo=spec["titulo"], anchos=[700, 260], ancho=1100)
            tmp = carpeta / (ident + "_tmp.png"); im.save(tmp)
            c = canvas.Canvas(str(ruta), pagesize=A4); w, h = A4
            c.drawImage(str(tmp), 30, h - 30 - (w - 60) * im.height / im.width, width=w - 60, height=(w - 60) * im.height / im.width); c.save(); tmp.unlink()
        ext = ".pdf"
    # decimales/unidades: marcados en el spec; el resto queda tal cual
    etiqueta = {"id": ident, "tipo": tipo, "archivo": ident + ext,
                "items": [{k: v for k, v in i.items()} for i in items], "no_inventario": sorted(set(t for t in no_inv if t and t.strip()))}
    (carpeta / (ident + ".json")).write_text(json.dumps(etiqueta, ensure_ascii=False, indent=1), encoding="utf-8")
    return etiqueta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--salida", default=str(Path(__file__).resolve().parents[2] / "datos-evaluacion" / "sinteticos"))
    args = ap.parse_args()
    carpeta = Path(args.salida); carpeta.mkdir(parents=True, exist_ok=True)
    specs = especificar()
    cuenta = {}
    for spec in specs:
        cuenta[spec["tipo"]] = cuenta.get(spec["tipo"], 0) + 1
        generar(spec, f"{spec['tipo']}_{cuenta[spec['tipo']]:02d}", carpeta)
    print(f"{len(specs)} documentos en {carpeta}: " + ", ".join(f"{k}={v}" for k, v in cuenta.items()))
    assert len(specs) == 39, len(specs)


if __name__ == "__main__":
    main()
